# -*- coding: utf-8 -*-
"""Per-task AI agent: propose an approach for a task (grounded in the lake schema
and precedent from earlier successful runs on similar tasks), then execute it -
run the data pulls, write the deliverable, attach it to the task as an HTML
report + comment, and suggest linked follow-up tasks."""
import datetime as dt
import json
import logging
import os
import re
import urllib.parse
import urllib.request

import tk_ai
import tk_db
import tk_skillbuilder
import lake_reader

LOG = logging.getLogger("tk_agent")

PLAN_SCHEMA = {
    "type": "object",
    "properties": {
        "approach": {"type": "string", "description": "2-4 sentence plan in plain English"},
        "deliverable_title": {"type": "string", "description": "short name for the output report"},
        "data_queries": {
            "type": "array", "minItems": 1, "maxItems": 4,
            "items": {
                "type": "object",
                "properties": {"label": {"type": "string"}, "sql": {"type": "string"}},
                "required": ["label", "sql"],
            },
        },
    },
    "required": ["approach", "deliverable_title", "data_queries"],
}

PLAN_SYSTEM = (
    "You are the task execution agent for Yo-Chi's finance TaskHub. Given a task and the "
    "data lake schema, design HOW to complete (or maximally advance) the task with data: "
    "an approach summary and 1-4 DuckDB queries producing the evidence. If precedents from "
    "similar past tasks are provided, prefer approaches that worked before, improved where "
    "obvious. SQL rules (STRICT): SELECT/WITH only; bare table names (never single-quoted "
    "FROM); CAST varchar dates AS DATE; aggregate to a useful grain; ORDER BY the metric "
    "that matters; <=200 rows per query. Data is refreshed nightly - say in the approach if "
    "true real-time data would be needed. Money is AUD net of GST. labour_cost in "
    "mart_venue_daily has >100%%-of-sales pay-run artefacts - exclude ratios above 1.0."
)

REPORT_SCHEMA = {
    "type": "object",
    "properties": {
        "summary": {"type": "string", "description": "3-6 sentence outcome summary for the task comment"},
        "report_html": {"type": "string", "description": "full deliverable as an HTML fragment (no doctype/html/body)"},
        "follow_up_tasks": {
            "type": "array", "maxItems": 4,
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "description": {"type": "string"},
                    "priority": {"type": "string", "enum": ["low", "medium", "high", "critical"]},
                },
                "required": ["title", "priority"],
            },
        },
    },
    "required": ["summary", "report_html"],
}

REPORT_SYSTEM = (
    "You are completing a finance task for Yo-Chi using the data pulls provided. Produce a "
    "sharp, decision-ready deliverable: a short summary (task comment thread) and report_html "
    "- an HTML FRAGMENT (no doctype/html/head/body) using only h2, h3, p, ul, ol, li, table, "
    "thead, tbody, tr, th, td, strong, and span with class rag-red / rag-amber / rag-green "
    "for status badges. REQUIRED structure, in order:\n"
    "  <h2>Executive summary</h2> - 3 sentences max + one RAG badge for overall status.\n"
    "  <h2>Key findings</h2> - the evidence, as compact TABLES (right numbers, right rows; "
    "top items only, never a raw dump), each with a one-line takeaway paragraph.\n"
    "  <h2>Analysis</h2> - what the numbers MEAN: causes, materiality, what is signal vs "
    "artefact (e.g. nightly-refresh lag), risks if unaddressed.\n"
    "  <h2>Proposal</h2> - a numbered action list: concrete steps, who should act (if "
    "inferable from the task), and by when relative to the close timetable.\n"
    "  <h2>Caveats</h2> - one short paragraph: data freshness and what a human must verify.\n"
    "Be quantitative and specific (venues, suppliers, dollars, dates). If the work surfaced "
    "distinct NEW issues deserving their own task, list up to 4 follow_up_tasks - only "
    "genuinely separate work. Fast Food Award items are 'areas to review', never 'breaches'. "
    "Money is AUD net of GST."
)


def _task_context(task_id):
    tasks = tk_db.get("tasks", {"id": "eq." + task_id, "select": "*"})
    if not tasks:
        raise ValueError("task not found")
    task = tasks[0]
    proj = tk_db.get("projects", {"id": "eq." + task["project_id"], "select": "name,description"})
    return task, (proj[0] if proj else {})


def _precedents(title):
    """Learning loop: (plans, lessons). Plans = recent successful plans for
    similar-sounding tasks. Lessons = user refinement feedback recorded on past
    successful runs (similar tasks first, then most recent anywhere) - the
    auto-refine loop: every 'Refine the output' instruction teaches later runs."""
    words = [w for w in re.findall(r"[A-Za-z]{4,}", title)][:4]
    runs = tk_db.get("agent_runs", {
        "outcome": "eq.success",
        "order": "created_at.desc",
        "limit": "40",
        "select": "task_title,plan,created_at",
    })
    scored = []
    for r in runs:
        score = sum(1 for w in words if w.lower() in r["task_title"].lower())
        if score:
            scored.append((score, r))
    scored.sort(key=lambda x: -x[0])
    precedents = [r for _, r in scored[:3]] if words else []
    matched = [r for _, r in scored]
    seen, lessons = set(), []
    for r in matched + [r for r in runs if r not in matched]:
        fb = (r.get("plan") or {}).get("feedback")
        if fb and fb not in seen:
            seen.add(fb)
            lessons.append(fb)
        if len(lessons) >= 5:
            break
    return precedents, lessons


def propose(task_id):
    task, proj = _task_context(task_id)
    schema = tk_skillbuilder._schema_text()
    precedents, lessons = _precedents(task["title"])
    prec_txt = ""
    if precedents:
        prec_txt = "\n\nPrecedents (successful plans for similar tasks):\n" + json.dumps(
            [{"task": p["task_title"], "plan": p["plan"]} for p in precedents],
            separators=(",", ":"))[:8000]
    if lessons:
        prec_txt += ("\n\nLessons from the user's past refinements of agent output "
                     "(apply proactively where relevant):\n" +
                     "\n".join("- " + l for l in lessons))
    user = "Lake schema:\n%s\n\nTask: %s\nProject: %s\nDue: %s\nDescription:\n%s%s" % (
        schema, task["title"], proj.get("name", ""), task.get("due_date"),
        task.get("description") or "(none)", prec_txt)
    plan = tk_ai.structured(PLAN_SYSTEM, user, "plan_task", PLAN_SCHEMA, max_tokens=4000)
    validation = tk_skillbuilder._validate(plan.get("data_queries") or [])
    if not validation or any(not v["ok"] for v in validation):
        repair = ("Task: %s\n\nOriginal approach (KEEP this as the approach text - do not "
                  "replace it with commentary about the fix):\n%s\n\nYour queries:\n%s\n\n"
                  "Validation errors:\n%s\n\nLake schema:\n%s\n\n"
                  "Return the corrected full plan WITH 1-4 concrete data_queries - "
                  "a plan with no queries is invalid.") % (
            task["title"], plan.get("approach", ""),
            json.dumps(plan.get("data_queries") or "(none returned)"),
            json.dumps(validation) or "(no queries)", schema)
        plan = tk_ai.structured(PLAN_SYSTEM, repair, "plan_task", PLAN_SCHEMA, max_tokens=4000)
        validation = tk_skillbuilder._validate(plan.get("data_queries") or [])
    plan["validation"] = validation
    plan["all_valid"] = bool(validation) and all(v["ok"] for v in validation)
    plan["used_precedents"] = [p["task_title"] for p in precedents]
    return plan


def _upload(task_id, filename, data, mime, uploader):
    url = "%s/storage/v1/object/taskapp-files/%s/%s" % (
        os.environ["SUPABASE_URL"].rstrip("/"), task_id, urllib.parse.quote(filename))
    req = urllib.request.Request(url, data=data, headers={
        "Authorization": "Bearer " + os.environ["SUPABASE_SERVICE_ROLE_KEY"],
        "apikey": os.environ["SUPABASE_SERVICE_ROLE_KEY"],
        "Content-Type": mime,
        "x-upsert": "true",
    }, method="POST")
    urllib.request.urlopen(req, timeout=60).read()
    tk_db.insert("attachments", [{
        "task_id": task_id,
        "storage_path": "%s/%s" % (task_id, filename),
        "filename": filename,
        "size_bytes": len(data),
        "mime": mime,
        "uploaded_by": uploader,
    }])


def _xlsx_bytes(pulls):
    """One sheet per data pull with the full result set."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for label, result in pulls.items():
        name = re.sub(r"[\\/*?:\[\]]", "-", label)[:31] or "data"
        ws = wb.create_sheet(name)
        if "error" in result:
            ws.append(["query error", result["error"]])
            continue
        ws.append(result["columns"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D683D")
        for row in result["rows"]:
            ws.append(["" if v is None else v for v in row])
        ws.freeze_panes = "A2"
        for i, col in enumerate(result["columns"], 1):
            width = max(len(str(col)) + 2, 12)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes(full_html):
    """Best-effort PDF of the report; returns None if conversion fails."""
    try:
        import io
        from xhtml2pdf import pisa
        buf = io.BytesIO()
        result = pisa.CreatePDF(full_html, dest=buf, encoding="utf-8")
        if result.err:
            return None
        return buf.getvalue()
    except Exception:
        LOG.exception("pdf conversion failed")
        return None


def _report_html(title, body_html, queries_used):
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;")
    q_html = "".join(
        "<details><summary>%s</summary><pre>%s</pre></details>" % (esc(q["label"]), esc(q["sql"]))
        for q in queries_used)
    return """<!DOCTYPE html><html><head><meta charset="utf-8"><title>%s</title>
<style>
  body{margin:0;background:#eef1f5;font-family:-apple-system,Segoe UI,Roboto,Arial,sans-serif;color:#1f2937}
  .report h2{font-size:16px;margin:22px 0 8px;color:#1d683d;border-bottom:2px solid #dbf2e4;padding-bottom:4px}
  .report h3{font-size:13.5px;margin:14px 0 6px;color:#374151}
  .report p,.report li{font-size:13.5px;line-height:1.55}
  .report table{border-collapse:collapse;width:100%%;margin:8px 0 12px;font-size:12.5px}
  .report th{background:#f0faf4;text-align:left;padding:6px 9px;border:1px solid #e3e7ee;font-weight:650}
  .report td{padding:5px 9px;border:1px solid #e3e7ee}
  .report tr:nth-child(even) td{background:#fafbfc}
  .rag-red,.rag-amber,.rag-green{display:inline-block;padding:2px 10px;border-radius:99px;font-size:11.5px;font-weight:700}
  .rag-red{background:#fee2e2;color:#b91c1c}.rag-amber{background:#fef3c7;color:#b45309}.rag-green{background:#dbf2e4;color:#1d683d}
  details{margin:3px 0}summary{cursor:pointer}pre{overflow-x:auto;background:#f6f8fa;padding:8px;border-radius:8px;font-size:10.5px}
</style></head>
<body>
<div style="max-width:820px;margin:0 auto;padding:24px">
  <div style="background:#1d683d;color:#fff;border-radius:12px 12px 0 0;padding:18px 24px">
    <div style="font-size:18px;font-weight:750">%s</div>
    <div style="font-size:12px;color:#c7e5d2;margin-top:2px">Generated by the TaskHub agent · %s · data as at last nightly lake refresh</div>
  </div>
  <div class="report" style="background:#fff;padding:6px 24px 18px;border:1px solid #e3e7ee;border-top:0">%s</div>
  <div style="background:#fff;border:1px solid #e3e7ee;border-top:0;border-radius:0 0 12px 12px;padding:12px 24px;font-size:11px;color:#6b7280">
    <div style="font-weight:700;margin-bottom:4px">Queries used</div>%s
  </div>
</div></body></html>""" % (esc(title), esc(title), dt.date.today().strftime("%d %b %Y"), body_html, q_html)


def download_file(task_id, filename):
    url = "%s/storage/v1/object/taskapp-files/%s/%s" % (
        os.environ["SUPABASE_URL"].rstrip("/"), task_id, urllib.parse.quote(filename))
    req = urllib.request.Request(url, headers={
        "Authorization": "Bearer " + os.environ["SUPABASE_SERVICE_ROLE_KEY"],
        "apikey": os.environ["SUPABASE_SERVICE_ROLE_KEY"]})
    return urllib.request.urlopen(req, timeout=60).read()


def execute(task_id, plan, requester_uid, feedback=None, prior_summary=None):
    task, _ = _task_context(task_id)
    queries = plan.get("data_queries") or []
    if not queries:
        raise ValueError("plan has no queries")
    referenced = set()
    for q in queries:
        referenced.update(re.findall(r"\b(?:FROM|JOIN)\s+\"?([A-Za-z_][A-Za-z0-9_]*)\"?",
                                     q["sql"], re.I))
    lake_reader.sync(extra_tables=referenced, log=LOG.info)
    pulls = {}
    for q in queries:
        try:
            pulls[q.get("label", "data")] = lake_reader.query(q["sql"], max_rows=200)
        except Exception as e:
            pulls[q.get("label", "data")] = {"error": str(e)[:200]}
    revision = ""
    if feedback:
        revision = ("\n\nTHIS IS A REVISION. Previous report summary:\n%s\n\n"
                    "The user's feedback to incorporate (follow it precisely):\n%s") % (
            prior_summary or "(not available)", feedback)
    lessons_txt = ""
    try:
        _, lessons = _precedents(task["title"])
        if lessons:
            lessons_txt = ("\n\nLessons from the user's past refinements of agent "
                           "reports (apply where relevant):\n" +
                           "\n".join("- " + l for l in lessons))
    except Exception:
        LOG.exception("lesson lookup failed (non-fatal)")
    user = "Task: %s\nDescription:\n%s\n\nApproach taken:\n%s\n\nData pulls:\n%s%s%s" % (
        task["title"], task.get("description") or "",
        plan.get("approach", ""), json.dumps(pulls, separators=(",", ":"))[:70000],
        lessons_txt, revision)
    try:
        out = tk_ai.structured(REPORT_SYSTEM, user, "deliver", REPORT_SCHEMA, max_tokens=8000)
        # a truncated generation can drop fields even with a forced tool
        body = out.get("report_html") or out.get("summary")
        if not body:
            raise RuntimeError("AI returned no report content (likely truncated) - try again")
        if not out.get("summary"):
            out["summary"] = body[:400]
    except Exception as e:
        tk_db.insert("agent_runs", [{
            "task_id": task_id, "task_title": task["title"], "plan": plan,
            "outcome": "error", "error": str(e)[:300], "requested_by": requester_uid}])
        raise

    title = plan.get("deliverable_title") or ("Agent report - " + task["title"][:60])
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", title)[:60] + "_" + dt.date.today().isoformat()
    full_html = _report_html(title, body, queries)

    files = []
    _upload(task_id, base + ".html", full_html.encode("utf-8"), "text/html", requester_uid)
    files.append(base + ".html")
    try:
        _upload(task_id, base + ".xlsx", _xlsx_bytes(pulls),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                requester_uid)
        files.append(base + ".xlsx")
    except Exception:
        LOG.exception("xlsx build failed")
    pdf = _pdf_bytes(full_html)
    if pdf:
        _upload(task_id, base + ".pdf", pdf, "application/pdf", requester_uid)
        files.append(base + ".pdf")

    tk_db.insert("comments", [{
        "task_id": task_id, "author_id": requester_uid,
        "body": "🤖 Agent run complete — attached: %s.\n\n%s" % (", ".join(files), out["summary"]),
    }])
    plan_record = {k: plan[k] for k in ("approach", "deliverable_title", "data_queries")
                   if k in plan}
    if feedback:
        # the refinement that shaped this successful run - future _precedents
        # calls surface it as a lesson
        plan_record["feedback"] = feedback[:600]
    tk_db.insert("agent_runs", [{
        "task_id": task_id, "task_title": task["title"],
        "plan": plan_record,
        "outcome": "success", "requested_by": requester_uid}])
    import tk_position
    tk_position.append_event(task["project_id"], "Agent %s '%s': %s" % (
        "revised" if feedback else "completed", task["title"], out["summary"][:250]))
    return {
        "summary": out["summary"],
        "attachment": ", ".join(files),
        "files": files,
        "follow_up_tasks": out.get("follow_up_tasks", []),
        "project_id": task["project_id"],
    }
