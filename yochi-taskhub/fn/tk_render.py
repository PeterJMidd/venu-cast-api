# -*- coding: utf-8 -*-
"""Render a routine's run - the AI review plus its data pulls - into the
formats a person asked to receive: csv, xlsx, pdf. One function per format,
all returning email-ready {name, data, mime} attachments.

Pulls arrive as lake_reader query results: {"columns": [...], "rows": [[...]]}.
A pull that errored arrives as {"error": "..."} and is shown as such rather
than dropped - an empty section that looks fine is worse than a named failure."""
import csv
import datetime as dt
import io
import logging
import re

LOG = logging.getLogger("tk_render")

MIME = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _safe(name, ext):
    stem = re.sub(r"[^A-Za-z0-9 _-]+", "", name or "report").strip() or "report"
    return "%s %s.%s" % (stem[:60], dt.date.today().isoformat(), ext)


def to_csv(name, pulls):
    """One CSV per data pull (there are at most a few). A single file per pull
    keeps each openable in Excel without guessing at section boundaries."""
    out = []
    for label, pull in (pulls or {}).items():
        buf = io.StringIO()
        w = csv.writer(buf)
        if pull.get("error"):
            w.writerow(["error"])
            w.writerow([pull["error"]])
        else:
            w.writerow(pull.get("columns") or [])
            for r in pull.get("rows") or []:
                w.writerow(["" if v is None else v for v in r])
        out.append({"name": _safe("%s - %s" % (name, label), "csv"),
                    "data": buf.getvalue().encode("utf-8-sig"),  # BOM: Excel opens UTF-8 cleanly
                    "mime": MIME["csv"]})
    return out


def to_xlsx(name, review, pulls):
    """One workbook: a Review sheet with the narrative, then a sheet per pull."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review"
    ws["A1"] = name
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Generated %s" % dt.datetime.now().strftime("%d %b %Y %H:%M")
    ws["A2"].font = Font(size=9, color="777777")
    row = 4
    for line in (review or "").splitlines():
        ws.cell(row=row, column=1, value=line)
        row += 1
    ws.column_dimensions["A"].width = 110

    for label, pull in (pulls or {}).items():
        sheet = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", " ", label)[:31] or "data")
        if pull.get("error"):
            sheet["A1"] = "This pull failed: %s" % pull["error"]
            continue
        cols = pull.get("columns") or []
        for j, c in enumerate(cols, 1):
            cell = sheet.cell(row=1, column=j, value=c)
            cell.font = Font(bold=True)
        for i, r in enumerate(pull.get("rows") or [], 2):
            for j, v in enumerate(r, 1):
                sheet.cell(row=i, column=j, value=v)
        for j, c in enumerate(cols, 1):
            sheet.column_dimensions[get_column_letter(j)].width = min(
                40, max(12, len(str(c)) + 4))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return [{"name": _safe(name, "xlsx"), "data": buf.getvalue(),
             "mime": MIME["xlsx"]}]


def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


PDF_ROW_CAP = 45   # a PDF is for reading; the full data rides in xlsx/csv


def to_pdf(name, review, pulls):
    from xhtml2pdf import pisa
    parts = [
        "<html><head><style>"
        "body{font-family:Helvetica,Arial,sans-serif;font-size:9pt;color:#222}"
        "h1{font-size:15pt;margin:0 0 2pt}"
        ".sub{font-size:8pt;color:#777;margin-bottom:12pt}"
        "pre{font-size:9pt;white-space:pre-wrap;font-family:Helvetica}"
        "h2{font-size:11pt;margin:14pt 0 4pt}"
        "table{border-collapse:collapse;width:100%;font-size:7.5pt}"
        "th{background:#f0f0f0;text-align:left;padding:3pt;border:0.5pt solid #ccc}"
        "td{padding:3pt;border:0.5pt solid #ddd}"
        ".note{font-size:7.5pt;color:#888}"
        "</style></head><body>",
        "<h1>%s</h1><div class='sub'>Generated %s &middot; Yo-Chi TaskHub</div>"
        % (_esc(name), dt.datetime.now().strftime("%d %b %Y %H:%M")),
        "<pre>%s</pre>" % _esc(review or "")]
    for label, pull in (pulls or {}).items():
        parts.append("<h2>%s</h2>" % _esc(label))
        if pull.get("error"):
            parts.append("<p class='note'>This pull failed: %s</p>" % _esc(pull["error"]))
            continue
        rows = pull.get("rows") or []
        parts.append("<table><tr>%s</tr>" % "".join(
            "<th>%s</th>" % _esc(c) for c in pull.get("columns") or []))
        for r in rows[:PDF_ROW_CAP]:
            parts.append("<tr>%s</tr>" % "".join(
                "<td>%s</td>" % _esc("" if v is None else v) for v in r))
        parts.append("</table>")
        if len(rows) > PDF_ROW_CAP:
            parts.append("<p class='note'>Showing %d of %d rows - the full data "
                         "is in the Excel/CSV attachment.</p>"
                         % (PDF_ROW_CAP, len(rows)))
    parts.append("</body></html>")
    buf = io.BytesIO()
    result = pisa.CreatePDF(io.StringIO("".join(parts)), dest=buf)
    if result.err:
        raise RuntimeError("PDF rendering failed")
    return [{"name": _safe(name, "pdf"), "data": buf.getvalue(),
             "mime": MIME["pdf"]}]


def render(formats, name, review, pulls):
    """Every requested format, in one list of attachments. An unknown format is
    reported, not ignored - the person asked for it by name."""
    out, problems = [], []
    for f in [x.strip().lower() for x in (formats or "").split(",") if x.strip()]:
        try:
            if f == "csv":
                out.extend(to_csv(name, pulls))
            elif f in ("xlsx", "excel"):
                out.extend(to_xlsx(name, review, pulls))
            elif f == "pdf":
                out.extend(to_pdf(name, review, pulls))
            else:
                problems.append("unknown format '%s'" % f)
        except Exception as e:
            LOG.exception("rendering %s failed", f)
            problems.append("%s failed: %s" % (f, str(e)[:120]))
    return out, problems
