import io
import logging

import openpyxl


def read_venue_configs(excel_bytes: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), read_only=True)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_map = {h.strip().lower(): i for i, h in enumerate(headers) if h}

    venue_col = header_map.get("venue")
    sender_col = header_map.get("senderemail")
    if venue_col is None or sender_col is None:
        raise ValueError(f"Excel must have 'Venue' and 'SenderEmail' columns. Found: {headers}")

    recipient_cols = []
    for i in range(1, 13):
        key = f"recipient{i}"
        if key in header_map:
            recipient_cols.append(header_map[key])

    configs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        venue = row[venue_col]
        if not venue:
            continue
        sender = row[sender_col]
        recipients = []
        for col_idx in recipient_cols:
            if col_idx < len(row) and row[col_idx]:
                recipients.append(str(row[col_idx]).strip())

        configs.append({
            "venue": str(venue).strip(),
            "sender_email": str(sender).strip(),
            "recipients": recipients,
        })
        logging.info(f"Config loaded: {venue} -> {len(recipients)} recipients")

    wb.close()
    return configs
